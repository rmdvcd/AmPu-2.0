"""XLSX/CSV/PNG/PDF y sesión reproducible con persistencia real."""

import json

import matplotlib
matplotlib.use("Agg", force=True)
import matplotlib.pyplot as plt
import pandas as pd
import pytest
from openpyxl import load_workbook

from ampu.io import ServicioExportacion as io


def test_excel_hojas_formatos_filas(tmp_path):
    a = pd.DataFrame({"x": [1, 2, 3], "P(X≤x)": [.1, .3, 1.]})
    b = pd.DataFrame({"Estrato": ["Norte", "Sur"], "nh": [3, 5]})
    archivo = io.excel({"Distribución": a, "Muestreo": b}, tmp_path / "datos")
    assert archivo.exists() and archivo.suffix == ".xlsx"
    libro = load_workbook(archivo)
    assert libro.sheetnames == ["Distribución", "Muestreo"]
    assert libro.worksheets[0].max_row == 4
    assert [c.value for c in libro.worksheets[0][1]] == ["x", "P(X≤x)"]
    assert libro.worksheets[0][4][1].value == 1.
    assert libro.worksheets[0][1][0].font.bold
    assert libro.worksheets[0].freeze_panes == "A2"
    assert libro.worksheets[1].max_row == 3
    csv = io.csv(a, tmp_path / "salida")
    assert csv.read_text(encoding="utf-8-sig").splitlines()[0] == "x,P(X≤x)"
    assert pd.read_csv(csv).shape == (3, 2)


def test_png_pdf_y_sesion(tmp_path):
    fig, ax = plt.subplots()
    ax.plot([0,1,2], [1,2,3])
    assert io.png(fig, tmp_path / "figura").stat().st_size > 500
    informe = io.pdf(tmp_path / "informe", "Informe AmPu", "Conclusión con ñ y á.",
                     {"Datos": pd.DataFrame({"x": [1,2], "Valor": [.4,.6]})}, fig, r"\frac{1}{2}")
    assert informe.read_bytes().startswith(b"%PDF-")
    assert informe.stat().st_size > 2_000
    plt.close(fig)
    estado = {"modulo": "distribuciones", "semilla": 42, "parametros": {"p": .3}, "datos": [1,2,3]}
    ruta = io.guardar(estado, tmp_path / "estado")
    assert json.loads(ruta.read_text(encoding="utf-8"))["version"] == 1
    assert io.cargar(ruta) == estado
    ruta.write_text('{"version": 99, "estado": {}}', encoding="utf-8")
    with pytest.raises(ValueError):
        io.cargar(ruta)
