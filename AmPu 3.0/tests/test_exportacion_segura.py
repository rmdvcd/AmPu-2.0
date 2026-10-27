"""Ediciones de tabla no deben convertirse en fórmulas activas al exportar."""

import pandas as pd
from openpyxl import load_workbook

from ampu.io.excel import exportar_csv, exportar_excel


def test_excel_y_csv_desactivan_formulas_pegadas(tmp_path):
    valores = ["=1+1", "+SUM(1,2)", "-SUM(1,2)", "@SUM(1,2)", "  =HYPERLINK(\"url\")", "\t=2+2", "texto", -42]
    tabla = pd.DataFrame({"=cabecera": valores, "numero": range(len(valores))})
    archivo = exportar_excel({"Sesión": tabla}, tmp_path / "sesion.xlsx")
    libro = load_workbook(archivo, read_only=True, data_only=False)
    hoja = libro.active
    assert hoja["A1"].value.startswith("'=")
    for fila in range(2, 8):
        assert hoja.cell(fila, 1).data_type == "s"
        assert hoja.cell(fila, 1).value.startswith("'")
    assert hoja["A8"].value == "texto"
    assert hoja["A9"].value == -42
    libro.close()

    csv = exportar_csv(tabla, tmp_path / "sesion.csv")
    texto = csv.read_text(encoding="utf-8-sig")
    assert texto.startswith("'=cabecera")
    assert "'=1+1" in texto and "'  =HYPERLINK" in texto
    assert "'-SUM" in texto and "'\t=2+2" in texto
    assert ",7\n" in texto
    assert tabla.iloc[0, 0] == "=1+1"  # La exportación no modifica la sesión.
