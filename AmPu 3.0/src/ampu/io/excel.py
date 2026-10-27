"""Exportación XLSX estilizada y CSV interoperable UTF-8."""

import re
from collections.abc import Mapping
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_DEF_AZUL = "172A4A"


def _texto_seguro(texto: str) -> str:
    """Evita que una celda pegada se interprete como fórmula al abrir CSV/XLSX."""
    peligroso = texto.startswith(("\t", "\r", "\n")) or texto.lstrip().startswith(("=", "+", "-", "@"))
    return "'" + texto if peligroso else texto


def _valor(v: object) -> object:
    """Evita tipos NumPy y valores NaN/inf que Excel no representa bien."""
    if v is None or (isinstance(v, (float, np.floating)) and np.isnan(v)):
        return None
    if isinstance(v, (float, np.floating)) and np.isinf(v):
        return "+∞" if v > 0 else "−∞"
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, (list, np.ndarray, tuple)):
        return _texto_seguro(", ".join(map(str, v)))
    return _texto_seguro(v) if isinstance(v, str) else v


def exportar_excel(tablas: Mapping[str, pd.DataFrame], destino: str | Path) -> Path:
    """Guarda una hoja por tabla, escribiendo el libro una sola vez al final.

    Returns:
        Ruta del archivo creado, con sufijo .xlsx.
    """
    if not tablas:
        raise ValueError("Debe haber al menos una tabla para exportar.")
    ruta = Path(destino).with_suffix(".xlsx")
    libro = Workbook()
    libro.remove(libro.active)
    usados: set[str] = set()
    for titulo, tabla in tablas.items():
        if not isinstance(tabla, pd.DataFrame):
            raise TypeError("Cada hoja necesita un DataFrame de pandas.")
        base = re.sub(r"[\\/*?:\[\]]", "_", str(titulo)).strip()[:31] or "Datos"
        nombre = base
        i = 2
        while nombre.lower() in usados:
            sufijo = f"_{i}"
            nombre = base[:31-len(sufijo)] + sufijo
            i += 1
        usados.add(nombre.lower())
        hoja = libro.create_sheet(nombre)
        hoja.append([_texto_seguro(str(nombre)) for nombre in tabla.columns])
        for fila in tabla.itertuples(index=False, name=None):
            hoja.append([_valor(v) for v in fila])
        for celda in hoja[1]:
            celda.fill = PatternFill("solid", fgColor=_DEF_AZUL)
            celda.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            celda.alignment = Alignment(horizontal="center")
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for columna in hoja.columns:
            letra = get_column_letter(columna[0].column)
            longitud = max((len(str(c.value or "")) for c in columna), default=12)
            hoja.column_dimensions[letra].width = min(55, max(13, longitud+3))
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, float):
                    celda.number_format = "0.0000000000"
    libro.save(ruta)
    return ruta


def exportar_csv(tabla: pd.DataFrame, destino: str | Path) -> Path:
    """Escribe encabezados y todas las filas con UTF-8-BOM para Excel."""
    if not isinstance(tabla, pd.DataFrame):
        raise TypeError("Se necesita una tabla de pandas.")
    ruta = Path(destino).with_suffix(".csv")
    limpia = tabla.copy()
    limpia.columns = [_texto_seguro(str(nombre)) for nombre in limpia.columns]
    for indice, tipo in enumerate(tabla.dtypes):
        if pd.api.types.is_object_dtype(tipo) or pd.api.types.is_string_dtype(tipo):
            limpia.iloc[:, indice] = tabla.iloc[:, indice].map(
                lambda valor: _texto_seguro(valor) if isinstance(valor, str) else valor)
    limpia.to_csv(ruta, index=False, encoding="utf-8-sig")
    return ruta
